"""Vuelca el contenido de una PPTX a Markdown, diapositiva por diapositiva.

Reconstruye el orden de lectura (agrupando cada título con su cuerpo), distingue
tablas, bloques monoespaciados y negritas, e incluye las notas del orador.
Útil para revisar o reutilizar el contenido sin abrir PowerPoint.

    python3 tools/dump_ppt_md.py Teoria-Lab-VLSM-DHCP.pptx > TEORIA-LAB-VLSM-DHCP.md
"""
import re
import sys
from pptx import Presentation
from pptx.util import Emu

FOOTER = "Redes de Computadoras · 8410"
MERGE = 0.45   # distancia vertical máxima para unir un título con su cuerpo
SAME_L = 0.06  # tolerancia horizontal para considerar dos cajas la misma columna


def para_md(p):
    """Devuelve (es_viñeta, markdown) reconstruyendo la negrita."""
    partes = []
    for r in p.runs:
        if not r.text:
            continue
        t = r.text.strip()
        if r.font.bold and t and not t.startswith(("▸", "–")):
            # conservar el espacio final del run para no pegar el texto siguiente
            partes.append("**%s**%s" % (t, " " if r.text.endswith(" ") else ""))
        else:
            partes.append(r.text)
    txt = "".join(partes).strip()
    # "**1.** texto" -> lista ordenada de markdown
    m = re.match(r"^\*\*(\d+)\.\*\*\s*(.*)$", txt)
    if m:
        return True, "%s. %s" % (m.group(1), m.group(2))
    if txt.startswith("▸"):
        return True, "- " + txt.lstrip("▸").strip()
    if txt.startswith("–"):
        return True, "  - " + txt.lstrip("–").strip()
    return False, txt


def es_mono(shape):
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
    return bool(runs) and all(r.font.name == "Consolas" for r in runs)


def tam_max(shape):
    sz = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
    return max(sz) if sz else 0


def forma_a_item(sh):
    """Normaliza una forma a (top, left, tipo, lineas) o None si se descarta."""
    if sh.has_table:
        t = sh.table
        filas = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in t.rows]
        if not filas:
            return None
        lineas = ["| " + " | ".join(c.replace("|", "\\|") for c in filas[0]) + " |",
                  "|" + "|".join(["---"] * len(filas[0])) + "|"]
        lineas += ["| " + " | ".join(c.replace("|", "\\|") for c in r) + " |"
                   for r in filas[1:]]
        return (round(Emu(sh.top).inches, 2), round(Emu(sh.left).inches, 2), "texto", lineas)
    if not sh.has_text_frame or not sh.text_frame.text.strip():
        return None
    crudo = sh.text_frame.text.strip()
    if crudo.startswith(FOOTER) or crudo == "→" or crudo.isdigit():
        return None
    out = [md for _, md in (para_md(p) for p in sh.text_frame.paragraphs) if md.strip()]
    if not out:
        return None
    if es_mono(sh):
        out = ["```text"] + out + ["```"]
    return (round(Emu(sh.top).inches, 2), round(Emu(sh.left).inches, 2), "texto", out)


def ordenar(items):
    """Une cada título con el bloque que tiene justo debajo en su columna."""
    items = sorted(items, key=lambda it: (it[0], it[1]))
    usados, grupos = set(), []
    for i, it in enumerate(items):
        if i in usados:
            continue
        lineas, top, left = list(it[3]), it[0], it[1]
        usados.add(i)
        for j, o in enumerate(items):
            if j in usados:
                continue
            if abs(o[1] - left) <= SAME_L and 0 < o[0] - top <= MERGE:
                lineas += o[3]
                top = o[0]
                usados.add(j)
                break
        grupos.append((top, left, lineas))
    grupos.sort(key=lambda g: (g[0], g[1]))
    return grupos


def titulo_de(items,Raw):
    """Separa título y subtítulo del resto del contenido de la diapositiva."""
    if not Raw:
        return "", "", []
    grande = [sh for sh in Raw if sh.has_text_frame and tam_max(sh) >= 24]
    if not grande:
        return "", "", []
    # el titulo es el texto mas grande; a igualdad, el mas alto
    t_sh = sorted(grande, key=lambda sh: (-tam_max(sh), Emu(sh.top).inches))[0]
    titulo = t_sh.text_frame.text.strip().replace("\n", " ")
    top_t = Emu(t_sh.top).inches
    subt, sub_sh, d_sub = "", None, 99
    for sh in Raw:
        if sh in grande or not sh.has_text_frame:
            continue
        d = abs(Emu(sh.top).inches - top_t)
        if d < 1.0 and 12 <= tam_max(sh) < 24:
            cand = sh.text_frame.text.strip().replace("\n", " ")
            if not subt or d < d_sub or (d == d_sub and len(cand) > len(subt)):
                subt, sub_sh, d_sub = cand, sh, d
    resto = [sh for sh in Raw if sh not in grande and sh is not sub_sh]
    return titulo, subt, resto


def dump(path):
    prs = Presentation(path)
    out, indice = [], []
    for i, s in enumerate(prs.slides, 1):
        Raw = [sh for sh in s.shapes
               if sh.has_table or (sh.has_text_frame and sh.text_frame.text.strip())]
        titulo, subt, resto = titulo_de(items=None, Raw=Raw)
        if titulo:
            indice.append((i, titulo))
        items = [it for it in (forma_a_item(sh) for sh in resto) if it]
        out.append("")
        out.append("## %d. %s" % (i, titulo or "Diapositiva %d" % i))
        if subt:
            out.append("")
            out.append("> %s" % subt)
        out.append("")
        for _, _, lineas in ordenar(items):
            out.extend(lineas)
            out.append("")
        notas = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        if notas:
            out.append("<details><summary>Notas del orador</summary>")
            out.append("")
            out.append(notas)
            out.append("")
            out.append("</details>")
        out.append("")
        out.append("---")
    return "\n".join(out), indice


if __name__ == "__main__":
    ruta = sys.argv[1] if len(sys.argv) > 1 else "Teoria-Lab-VLSM-DHCP.pptx"
    cuerpo, indice = dump(ruta)
    print(cuerpo)