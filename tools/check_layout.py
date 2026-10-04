"""Verificación de layout: estima si el texto cabe en su caja y detecta desbordes
hacia fuera de la diapositiva. Útil porque no podemos ver el render."""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
# ancho medio de carácter en pulgadas, por tamaño de fuente Calibri
CHAR_W = {10: 0.058, 11: 0.062, 12: 0.068, 13: 0.073, 14: 0.079, 15: 0.085,
          16: 0.090, 17: 0.095, 18: 0.101, 19: 0.107, 20: 0.113, 26: 0.147, 40: 0.226}
LINE_H = 1.22  # alto de línea como múltiplo del tamaño


def est_height(shape):
    """Altura estimada del texto, en pulgadas."""
    tf = shape.text_frame
    w_in = Emu(shape.width).inches - 0.1
    total = 0.0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not txt:
            total += 0.14
            continue
        sz = max((r.font.size.pt for r in p.runs if r.font.size), default=12)
        cw = CHAR_W.get(int(round(sz)), 0.068)
        chars_per_line = max(1, int(w_in / cw))
        lines = max(1, -(-len(txt) // chars_per_line))
        total += lines * sz * LINE_H / 72
        if p.space_after:
            total += p.space_after.pt / 72
    return total


def check(path):
    prs = Presentation(path)
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    problems = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            l, t = Emu(sh.left).inches, Emu(sh.top).inches
            w, h = Emu(sh.width).inches, Emu(sh.height).inches
            if l < -0.01 or t < -0.01 or l + w > W + 0.01 or t + h > H + 0.01:
                problems.append(f"  slide {i}: {sh.shape_type} FUERA DE LIENZO "
                                f"({l:.2f},{t:.2f},{w:.2f}x{h:.2f})")
            if sh.has_text_frame and sh.text_frame.text.strip():
                eh = est_height(sh)
                if eh > h + 0.12:
                    preview = sh.text_frame.text.strip().replace("\n", " ")[:60]
                    problems.append(f"  slide {i}: TEXTO SE PASA "
                                    f"(caja {h:.2f}in, necesita ~{eh:.2f}in) :: {preview!r}")
            if sh.has_table:
                tbl = sh.table
                th = sum(Emu(r.height).inches for r in tbl.rows)
                if t + th > H - 0.35:
                    problems.append(f"  slide {i}: TABLA se sale por abajo "
                                    f"(termina en {t+th:.2f}in)")
                for ri, row in enumerate(tbl.rows):
                    for ci, cell in enumerate(row.cells):
                        cw = Emu(tbl.columns[ci].width).inches
                        rt = cell.text_frame.text
                        rs = max((r.font.size.pt for p in cell.text_frame.paragraphs
                                  for r in p.runs if r.font.size), default=12)
                        cpl = max(1, int((cw - 0.13) / CHAR_W.get(int(round(rs)), 0.068)))
                        need = max(1, -(-len(rt) // cpl))
                        if need > 1 and len(rt) > 28:
                            problems.append(f"  slide {i}: celda ({ri},{ci}) envuelve "
                                            f"{need} líneas: {rt[:45]!r}")
    return len(prs.slides), problems


if __name__ == "__main__":
    for path in sys.argv[1:]:
        n, probs = check(path)
        print(f"\n=== {path} · {n} diapositivas ===")
        if not probs:
            print("  sin problemas de layout detectados")
        for p in probs:
            print(p)
